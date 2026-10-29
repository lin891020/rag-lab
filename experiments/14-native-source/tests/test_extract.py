"""html route: the law-system page stripped to text, one paragraph per line, chrome removed;
attachments: odt / docx / pdf become text (tables as 表頭：值｜… lines), .doc is reported as
unavailable."""

from pathlib import Path

import pytest

from extract import _shingles, attachment_text, html_text, is_novel

PAGE = """<html><head><title>教育部主管法規共用系統-法規內容</title></head><body>
<div id="header">跳到主要內容區塊 法規功能按鈕區 法規內容 條文檢索 法規沿革 歷史法規</div>
<div class="law-reg-content">
<div><span class="csp-s1">一、目的：教育部（以下簡稱本部）為協助各師資培育之大學，特訂定本要點。</span></div>
<div><span>二、補助對象：各師培大學。</span></div>
<p class="MsoNormal"><span>三、補助原則：</span></p>
<p>（一）本要點以部分補助為原則，<b>補助比率不得超過百分之九十</b>。</p>
<table><tr><td>項目</td><td>金額</td></tr><tr><td>上限</td><td>三百五十萬元</td></tr></table>
</div>
<div id="footer">本網站使用 IE11 瀏覽 版權所有</div>
</body></html>"""


def test_html_text_keeps_paragraphs_and_drops_chrome():
    text = html_text(PAGE)
    lines = text.splitlines()
    assert lines[0].startswith("一、目的：教育部") and "特訂定本要點。" in lines[0]
    assert "二、補助對象：各師培大學。" in lines and "三、補助原則：" in lines
    assert "補助比率不得超過百分之九十" in text  # inline tags do not break a line
    assert "法規功能按鈕區" not in text and "版權所有" not in text and "<" not in text
    assert "項目：上限｜金額：三百五十萬元" in text  # tables as key-value rows


def test_attachment_text_reads_odt_and_docx_and_refuses_doc(tmp_path: Path):
    pytest.importorskip("docx")
    pytest.importorskip("odf")
    from docx import Document
    from odf.opendocument import OpenDocumentText
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    d = Document()
    d.add_paragraph("附表一 收費基準")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "項目", "收費"
    t.cell(1, 0).text, t.cell(1, 1).text = "禮廳", "五千元"
    d.save(tmp_path / "a.docx")
    text, ok = attachment_text(tmp_path / "a.docx")
    assert ok and "附表一 收費基準" in text and "項目：禮廳｜收費：五千元" in text

    o = OpenDocumentText()
    o.text.addElement(P(text="附表二 說明"))
    tb = Table()
    for cells in (("名稱", "數量"), ("委員", "七人")):
        row = TableRow()
        for c in cells:
            cell = TableCell()
            cell.addElement(P(text=c))
            row.addElement(cell)
        tb.addElement(row)
    o.text.addElement(tb)
    o.save(str(tmp_path / "b.odt"))
    text, ok = attachment_text(tmp_path / "b.odt")
    assert ok and "附表二 說明" in text and "名稱：委員｜數量：七人" in text

    (tmp_path / "c.doc").write_bytes(b"\xd0\xcf\x11\xe0" + b"\0" * 100)
    text, ok = attachment_text(tmp_path / "c.doc")
    assert not ok and text == ""


@pytest.mark.parametrize("ext", ["xls", "zip"])
def test_unknown_attachment_types_are_unavailable(tmp_path: Path, ext: str):
    p = tmp_path / f"x.{ext}"
    p.write_bytes(b"xx")
    assert attachment_text(p) == ("", False)


def test_is_novel_rejects_an_attachment_that_repeats_the_body():
    body = "一、目的：" + "甲乙丙丁" * 60 + "\n二、對象：" + "戊己庚辛" * 60
    seen = _shingles(body)
    assert not is_novel(body + "\n（請參見附件）", seen)  # the gazette PDF attached again
    assert is_novel("附表一 收費基準\n" + "項目：禮廳｜收費：五千元\n" * 30, seen)
