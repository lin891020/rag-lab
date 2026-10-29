"""Three routes to text for the same notice: the gazette PDF (13's / 09's `ocr` text, reused),
the law-system HTML stripped to text (Tika-style: tags off, one paragraph per line, tables as
表頭：值｜… rows), and that text plus the attachments converted natively.

  pdf       corpus text of 13 (or 09) for the same sysid, copied
  html      corpus/native/<sysid>.html -> text
  html+att  html text + every attachment that converts (.odt via odfpy, .docx via python-docx,
            .pdf via pypdf) and is not a repeat of text already kept (the systems attach the gazette
            PDF itself and older versions); .doc (binary Word) and the rest are unavailable

    uv run --group office python experiments/14-native-source/extract.py
"""

from __future__ import annotations

import html as htmllib
import json
import logging
import re
import sys
from pathlib import Path

from rag_lab.log import setup_logging

HERE = Path(__file__).parent
CORPUS = HERE / "corpus"
NATIVE = CORPUS / "native"
SOURCES = {
    "13": HERE.parent / "13-points-tables" / "corpus" / "text",
    "09": HERE.parent / "09-messy-real" / "corpus" / "text",
}
logger = logging.getLogger(__name__)
_QUOTED = re.compile(r"「([^」]+)」")
_BARE = re.compile(
    r"(?:增訂|修正|訂定|廢止)([一-鿿]+?(?:自治條例|條例|法|規則|辦法|標準|準則|要點|規程|細則))"
)


def law_name(caption: str) -> str:
    """09's: the first 「quoted」 name, else the name after 修正／增訂 (laws passed by the
    legislature are captioned without quotes)."""
    m = _QUOTED.search(caption) or _BARE.search(caption)
    return m.group(1) if m else ""


_BLOCK = re.compile(r"</?(?:p|div|br|li|h[1-6]|tr|section|article)\b[^>]*>", re.I)
_TAG = re.compile(r"<[^>]+>")
_SCRIPT = re.compile(r"<(script|style|noscript)\b.*?</\1>", re.S | re.I)
_TABLE = re.compile(r"<table\b.*?</table>", re.S | re.I)
_ROW = re.compile(r"<tr\b.*?</tr>", re.S | re.I)
_CELL = re.compile(r"<t[dh]\b[^>]*>(.*?)</t[dh]>", re.S | re.I)
_CHROME = (
    "跳到主要內容區塊",
    "法規功能按鈕區",
    "法規內容 條文檢索",
    "您的瀏覽器不支援",
    "版權所有",
    "列印時間",
    "瀏覽",
    "Download",
)


def _clean(s: str) -> str:
    return re.sub(r"[ \t　]+", " ", htmllib.unescape(_TAG.sub("", s))).strip()


def _table_rows(table_html: str) -> list[str]:
    rows = []
    for r in _ROW.findall(table_html):
        cells = [_clean(c) for c in _CELL.findall(r)]
        if any(cells):
            rows.append(cells)
    if len(rows) < 2:
        return [" ".join(c for c in rows[0] if c)] if rows else []
    header: list[str] = []
    for h in rows[0]:
        header.append(h or (header[-1] if header else ""))
    out = []
    for r in rows[1:]:
        pairs = [f"{h}：{v}" for h, v in zip(header, r, strict=False) if v]
        if pairs:
            out.append("｜".join(pairs))
    return out


def _body(page: str) -> str:
    """The law content region when the page marks one; else the whole body."""
    m = re.search(r'<div[^>]+class="[^"]*law-reg-content[^"]*"[^>]*>(.*)', page, re.S | re.I)
    if m:
        return m.group(1)
    m = re.search(r"<body\b[^>]*>(.*)</body>", page, re.S | re.I)
    return m.group(1) if m else page


def html_text(page: str) -> str:
    """Tags stripped, block tags as line breaks, tables as key-value rows, site chrome dropped."""
    page = _SCRIPT.sub(" ", page)
    body = _body(page)
    body = _TABLE.sub(lambda m: "\n" + "\n".join(_table_rows(m.group(0))) + "\n", body)
    body = _BLOCK.sub("\n", body)
    lines = []
    for ln in htmllib.unescape(_TAG.sub("", body)).splitlines():
        ln = re.sub(r"[ \t　]+", " ", ln).strip()
        if not ln or any(k in ln for k in _CHROME):
            continue
        lines.append(ln)
    return "\n".join(lines)


def _docx_text(path: Path) -> str:
    from docx import Document  # the optional `office` group

    d = Document(str(path))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        rows = [r for r in rows if any(r)]
        if len(rows) >= 2:
            header: list[str] = []
            for h in rows[0]:
                header.append(h or (header[-1] if header else ""))
            for r in rows[1:]:
                pairs = [f"{h}：{v}" for h, v in zip(header, r, strict=False) if v]
                if pairs:
                    out.append("｜".join(pairs))
    return "\n".join(out)


def _odt_text(path: Path) -> str:
    from odf import table, teletype, text
    from odf.opendocument import load

    doc = load(str(path))
    out = []
    tables = {id(t) for t in doc.getElementsByType(table.Table)}
    for node in doc.text.childNodes:
        if node.qname[1] == "table":
            rows = []
            for r in node.getElementsByType(table.TableRow):
                cells = [
                    teletype.extractText(c).strip() for c in r.getElementsByType(table.TableCell)
                ]
                if any(cells):
                    rows.append(cells)
            if len(rows) >= 2:
                header: list[str] = []
                for h in rows[0]:
                    header.append(h or (header[-1] if header else ""))
                for r in rows[1:]:
                    pairs = [f"{h}：{v}" for h, v in zip(header, r, strict=False) if v]
                    if pairs:
                        out.append("｜".join(pairs))
            continue
        if id(node) in tables:
            continue
        s = teletype.extractText(node).strip() if node.nodeType == node.ELEMENT_NODE else ""
        if s:
            out.append(s)
    _ = text  # keep the import for readers; odf.text types are matched by qname above
    return "\n".join(out)


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)


def _shingles(text: str, k: int = 30) -> set[str]:
    t = re.sub(r"\s+", "", text)
    return {t[i : i + k] for i in range(0, max(1, len(t) - k + 1), 7)}


def is_novel(text: str, seen: set[str], *, k: int = 30, threshold: float = 0.5) -> bool:
    """An attachment that mostly repeats text already in the corpus (the gazette PDF attached
    again, an older version of the same table) is not added, else the gold would be indexed
    several times and every hit rate would inflate."""
    mine = _shingles(text, k)
    if not mine:
        return False
    return sum(1 for sh in mine if sh in seen) / len(mine) < threshold


def attachment_text(path: Path | str) -> tuple[str, bool]:
    """(text, available). .doc and unknown types are unavailable without LibreOffice."""
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            return _docx_text(path), True
        if ext == ".odt":
            return _odt_text(path), True
        if ext == ".pdf":
            return _pdf_text(path), True
    except (OSError, ValueError, KeyError) as exc:
        logger.warning("attachment unreadable", extra={"ctx": {"path": path.name}}, exc_info=exc)
        return "", False
    return "", False


def main() -> int:
    setup_logging("rag_lab")
    docs = json.loads((CORPUS / "manifest.json").read_text("utf-8"))["docs"]
    (CORPUS / "text").mkdir(exist_ok=True)
    stats = {}
    for d in docs:
        sysid = d["sysid"]
        pdf_src = SOURCES[d["from"]] / f"{sysid}.ocr.txt"
        texts = {"pdf": pdf_src.read_text("utf-8")}
        texts["html"] = html_text((NATIVE / f"{sysid}.html").read_text("utf-8"))
        parts, att = [texts["html"]], []
        # repeats of the page body OR of the gazette PDF are not new content (the systems attach
        # the gazette itself and the 修正對照表 that restates the amended text)
        seen = _shingles(texts["html"]) | _shingles(texts["pdf"])
        for a in d.get("attachments", []):
            text, ok = attachment_text(NATIVE / a["file"])
            novel = ok and bool(text.strip()) and is_novel(text, seen)
            att.append({**a, "available": ok, "chars": len(text), "kept": novel})
            if novel:
                parts.append(f"【附件 {a['name']}】\n{text}")
                seen |= _shingles(text)
        texts["html+att"] = "\n".join(parts)
        for route, text in texts.items():
            (CORPUS / "text" / f"{sysid}.{route}.txt").write_text(text, "utf-8")
        stats[sysid] = {"chars": {k: len(v) for k, v in texts.items()}, "attachments": att}
    (CORPUS / "extract_stats.json").write_text(
        json.dumps(stats, ensure_ascii=False, indent=1), "utf-8"
    )
    logger.info("extracted", extra={"ctx": {"docs": len(docs)}})
    return 0


if __name__ == "__main__":
    sys.exit(main())
